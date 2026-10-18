"""Resume/JD text extraction: PDF, DOCX, DOC; Gemini OCR fallback for scanned PDFs."""
import hashlib
import re
import subprocess
import tempfile
from pathlib import Path

from . import store

MIN_WORDS = 50


def file_hash(path):
    return hashlib.sha1(Path(path).read_bytes()).hexdigest()


def _pdf_text(path):
    out = subprocess.run(["pdftotext", str(path), "-"], capture_output=True, timeout=60)
    return out.stdout.decode("utf-8", "ignore")


def _pdf_pages(path):
    out = subprocess.run(["pdfinfo", str(path)], capture_output=True, timeout=30)
    m = re.search(r"Pages:\s+(\d+)", out.stdout.decode("utf-8", "ignore"))
    return int(m.group(1)) if m else 0


def _docx_text(path):
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _doc_text(path):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "txt:Text", "--outdir", tmp, str(path)],
                       capture_output=True, timeout=120)
        out = Path(tmp) / (Path(path).stem + ".txt")
        return out.read_text("utf-8", "ignore") if out.exists() else ""


def clean(text):
    text = text.replace("\x0c", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract(path, allow_ocr=True):
    """Return {text, pages, words, ocr, error}. Cached by content hash."""
    path = Path(path)
    key = "text:" + file_hash(path)
    cached = store.get(key)
    if cached and not (allow_ocr and cached["words"] < MIN_WORDS and not cached.get("ocr_tried")):
        return cached
    ext = path.suffix.lower()
    res = {"text": "", "pages": 0, "ocr": False, "error": None}
    try:
        if ext == ".pdf":
            res["text"] = _pdf_text(path)
            res["pages"] = _pdf_pages(path)
        elif ext == ".docx":
            res["text"] = _docx_text(path)
            res["pages"] = 1
        elif ext == ".doc":
            res["text"] = _doc_text(path)
            res["pages"] = 1
        else:
            res["error"] = f"unsupported type {ext}"
    except Exception as e:  # corrupt files must not kill a batch
        res["error"] = f"{type(e).__name__}: {e}"
    res["text"] = clean(res["text"])
    res["words"] = len(res["text"].split())
    if allow_ocr and ext == ".pdf" and res["words"] < MIN_WORDS:
        res["ocr_tried"] = True
        try:
            from . import llm
            if llm.available():
                txt = llm.ocr_pdf(path.read_bytes())
                if txt and len(txt.split()) > res["words"]:
                    res["text"], res["ocr"] = clean(txt), True
                    res["words"] = len(res["text"].split())
        except Exception as e:
            res["error"] = f"ocr failed: {type(e).__name__}: {e}"
    store.put(key, res)
    return res
