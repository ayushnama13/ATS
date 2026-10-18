"""Excel export of a ranking, including the per-factor breakdown."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .features import FACTORS


def to_xlsx(rows, company, weights, enabled):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"
    head = ["Rank", "Roll", "Name", "Program", "Department", "Specialization", "CGPA", "Score"] + \
           [f"{f} (pts)" for f in FACTORS] + ["Matched must-have", "Missing must-have", "LLM verdict", "Format flags"]
    ws.append(head)
    for r in rows:
        d = r["detail"]
        rub = d.get("rubric") or {}
        ws.append([r["rank"], r["roll"], r["name"], r["program"], r["dept"], r["spec"], r["cgpa"], r["score"]] +
                  [r["parts"][f]["points"] if r["parts"][f]["weight"] else None for f in FACTORS] +
                  [", ".join(d["matched_must"]), ", ".join(d["missing_must"]), rub.get("summary", ""),
                   "; ".join(d["format_flags"])])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "D2"
    for i, w in enumerate([6, 12, 26, 14, 28, 26, 7, 8] + [9] * len(FACTORS) + [30, 30, 60, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions
    cfg = wb.create_sheet("Config")
    cfg.append(["Company", company])
    cfg.append(["Factor", "Enabled", "Weight"])
    for f in FACTORS:
        cfg.append([f, bool(enabled.get(f, True)), weights.get(f, 0)])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
